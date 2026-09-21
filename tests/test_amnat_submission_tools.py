from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import RGBColor

from scripts.build_anonymous_review_bundle import build
from scripts.format_amnat_review_docx import format_document
from scripts.verify_amnat_claims import verify


def test_registered_slk_claims_recompute() -> None:
    receipt = verify()
    assert receipt["all_checks_pass"] is True
    inv = receipt["checks"]["INV1_fixation_occupancy_invariant"]
    assert inv["comparisons"] == 336
    assert inv["derived_from_moran_process"] is True
    assert inv["derived_from_rare_mutation_chain"] is True
    assert inv["max_relative_error"] < 1e-10
    assert receipt["checks"]["CANONICAL_MAPPING_GUARD"]["guard_raised"] is True


def test_anonymous_bundle_is_curated_and_scanned(tmp_path: Path) -> None:
    out = build(tmp_path / "reviewer_bundle")
    assert (out / "README.md").is_file()
    assert (out / "MANUSCRIPT_SOURCE.md").is_file()
    assert (out / "CLAIM_VERIFICATION_RECEIPT.json").is_file()
    assert (out / "scripts" / "slk_threshold_atlas.py").is_file()
    assert (out / "tests" / "test_moran_process_invariant.py").is_file()
    assert (out / "docs" / "INV1_EXECUTABLE_VALIDATION_V1.md").is_file()
    assert (out / "ANONYMITY_AUDIT.txt").read_text(encoding="utf-8").startswith("identity_scan=PASS")
    assert (out / "SHA256SUMS.txt").is_file()
    assert not (out / ".git").exists()


def test_formatter_adds_line_and_page_number_fields() -> None:
    doc = Document()
    doc.add_heading("Anonymous review manuscript", level=1)
    doc.add_paragraph("A test paragraph.")
    format_document(doc, line_numbers=True)

    sect_pr = doc.sections[0]._sectPr
    assert sect_pr.find(qn("w:lnNumType")) is not None
    footer_xml = doc.sections[0].footer._element.xml
    assert " PAGE " in footer_xml
    assert "w:suppressLineNumbers" in footer_xml
    assert doc.styles["Normal"].paragraph_format.line_spacing == 2
    assert doc.styles["Heading 1"].font.color.rgb == RGBColor(0, 0, 0)
